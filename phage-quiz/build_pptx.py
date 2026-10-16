#!/usr/bin/env python3
"""Build phage-quiz.pptx — a standalone PowerPoint deck (runs on any PC with
PowerPoint / LibreOffice) mirroring the ICBRAR slide format:
question slide followed by an answer slide with the correct hexagon highlighted.
Correct answers are ALSO in the speaker notes of every question slide.
Figures for the visual round are rendered to PNG with PyMuPDF and embedded.

Usage: python3 build_pptx.py
"""
import json, re, pathlib
import pymupdf
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = pathlib.Path(__file__).parent
FIGDIR = HERE / "figs"
FIGDIR.mkdir(exist_ok=True)
data = json.loads(re.search(r"/\*QUIZ_JSON_START\*/(.*)/\*QUIZ_JSON_END\*/",
                            (HERE / "quiz-data.js").read_text(encoding="utf-8"), re.S).group(1))

BG = RGBColor(0x05, 0x07, 0x0D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUT = RGBColor(0x9A, 0xA7, 0xBD)
ACC = RGBColor(0xBD, 0x0D, 0xFF)
PINK = RGBColor(0xFF, 0x4D, 0xF0)
CARD = RGBColor(0x0C, 0x0F, 0x16)
LINE = RGBColor(0x3A, 0x43, 0x56)
G1 = RGBColor(0x0A, 0x66, 0xFF)
G2 = RGBColor(0x8A, 0x2B, 0xF5)
SKY = RGBColor(0x7D, 0xD3, 0xFC)

PT = 1 / 72.0  # pdf points -> inches

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s

def tb(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """runs: list of paragraphs; each paragraph = list of (text, size, color, bold)"""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for (t, size, color, bold) in para:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = "Arial"
    return box

def hexcard(s, x, y, w, h, letter, text, win=False, size=15):
    shp = s.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if win:
        f = shp.fill
        f.gradient()
        stops = f.gradient_stops
        stops[0].color.rgb = G1
        stops[1].color.rgb = G2
        try:
            f.gradient_angle = 0.0
        except Exception:
            pass
        shp.line.color.rgb = RGBColor(0x9C, 0x8C, 0xFF)
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = CARD
        shp.line.color.rgb = LINE
    shp.line.width = Pt(1)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.55)
    tf.margin_right = Inches(0.45)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = letter + "   "
    r.font.size = Pt(24); r.font.bold = True
    r.font.color.rgb = PINK if win else ACC
    r.font.name = "Arial"
    r2 = p.add_run(); r2.text = text
    r2.font.size = Pt(size); r2.font.bold = True
    r2.font.color.rgb = WHITE
    r2.font.name = "Arial"
    return shp

def chip(s, x, y, text, size=13):
    w = 0.16 * len(text) + 0.5
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.42))
    shp.shadow.inherit = False
    shp.fill.solid(); shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = LINE
    tf = shp.text_frame; tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = SKY; r.font.name = "Arial"
    return w

# ------------------------- figures (PyMuPDF -> PNG) -------------------------
FGCOL = (0.02, 0.027, 0.051)
WH = (1, 1, 1)
BL = (0.56, 0.72, 1.0)
MG = (0.741, 0.055, 1.0)
LAWN = (0.85, 0.79, 0.64)
CLEAR = (0.97, 0.95, 0.91)

def figpage(w, h):
    d = pymupdf.open()
    p = d.new_page(width=w, height=h)
    p.draw_rect(pymupdf.Rect(0, 0, w, h), fill=FGCOL, color=None)
    return d, p

def figtext(p, x, y, s, size=13, color=(0.81, 0.88, 1.0)):
    w = pymupdf.get_text_length(s, fontname="hebo", fontsize=size)
    p.insert_text((x - w / 2, y), s, fontname="hebo", fontsize=size, color=color)

def arrow(p, x0, y0, x1, y1):
    p.draw_line((x0, y0), (x1, y1), color=MG, width=3)
    import math
    ang = math.atan2(y1 - y0, x1 - x0)
    for da in (2.6, -2.6):
        p.draw_line((x1, y1), (x1 + 10 * math.cos(ang + da), y1 + 10 * math.sin(ang + da)), color=MG, width=3)

def savefig(d, name):
    pix = d[0].get_pixmap(dpi=200)
    fn = FIGDIR / f"{name}.png"
    pix.save(str(fn))
    d.close()
    return fn

def make_figures():
    # growth curve
    d, p = figpage(460, 280)
    p.draw_line((50, 30), (50, 230), color=WH, width=3)
    p.draw_line((50, 230), (430, 230), color=WH, width=3)
    pts = [(50, 195), (150, 195), (210, 190)]
    pts += [(212 + t * 4, 190 - (t ** 1.8) * 130) for t in range(0, 15)]
    pts += [(320, 58), (400, 58)]
    p.draw_polyline(pts, color=BL, width=4)
    p.draw_line((150, 195), (150, 230), color=(0.23, 0.26, 0.34), width=2, dashes="4 4")
    p.draw_line((212, 190), (212, 230), color=(0.23, 0.26, 0.34), width=2, dashes="4 4")
    arrow(p, 120, 150, 145, 182)
    figtext(p, 105, 142, "?", 26, MG)
    figtext(p, 240, 268, "Time (min)", 15)
    figtext(p, 95, 250, "adsorption", 12, (0.49, 0.55, 0.63))
    figtext(p, 181, 250, "flat phase", 12, (0.49, 0.55, 0.63))
    figtext(p, 330, 250, "rise -> plateau", 12, (0.49, 0.55, 0.63))
    savefig(d, "growth_curve")
    # plaque halo
    d, p = figpage(380, 300)
    p.draw_circle((190, 150), 130, fill=(0.06, 0.08, 0.11), color=BL, width=4)
    p.draw_circle((190, 150), 118, fill=LAWN, color=None)
    p.draw_circle((190, 150), 52, fill=CLEAR, color=None)
    p.draw_circle((190, 150), 26, fill=(1, 1, 1), color=None)
    arrow(p, 300, 60, 245, 112)
    figtext(p, 312, 54, "?", 26, MG)
    figtext(p, 190, 292, "bacterial lawn with plaque + outer clearing", 13, (0.49, 0.55, 0.63))
    savefig(d, "plaque_halo")
    # T4
    d, p = figpage(300, 330)
    p.draw_polyline([(150, 12), (205, 42), (205, 102), (150, 132), (95, 102), (95, 42)], color=WH, width=4, closePath=True)
    p.draw_rect(pymupdf.Rect(136, 132, 164, 218), color=WH, width=4)
    for y in (146, 160, 174, 188, 202):
        p.draw_line((136, y), (164, y), color=WH, width=2.5)
    p.draw_polyline([(150, 218), (178, 232), (150, 246), (122, 232)], color=WH, width=4, closePath=True)
    for a, b, c, e in [(128, 240, 48, 300), (172, 240, 252, 300), (136, 236, 112, 306), (164, 236, 188, 306)]:
        p.draw_bezier((a, b), (a - 20 if a < 150 else a + 20, b + 30), (c + 10 if c < 150 else c - 10, e - 20), (c, e), color=WH, width=3.5)
    arrow(p, 272, 250, 232, 276)
    figtext(p, 284, 246, "?", 26, MG)
    savefig(d, "t4_diagram")
    # lysogeny
    d, p = figpage(430, 240)
    p.draw_oval(pymupdf.Rect(30, 40, 270, 200), color=WH, width=4)
    p.draw_line((105, 108), (155, 108), color=MG, width=10)
    p.draw_line((60, 200), (110, 165), color=(0.23, 0.26, 0.34), width=2.5, dashes="4 4")
    p.draw_line((240, 200), (195, 165), color=(0.23, 0.26, 0.34), width=2.5, dashes="4 4")
    p.draw_oval(pymupdf.Rect(260, 124, 400, 216), color=BL, width=3)
    p.draw_line((305, 162), (331, 162), color=MG, width=7)
    arrow(p, 130, 96, 250, 60)
    figtext(p, 300, 52, "?", 26, MG)
    figtext(p, 215, 228, "host chromosome (loop) with inserted segment, copied to daughters", 12, (0.49, 0.55, 0.63))
    savefig(d, "lysogeny")
    # overlay
    d, p = figpage(460, 250)
    p.draw_rect(pymupdf.Rect(40, 40, 86, 190), color=WH, width=4, radius=0.2)
    p.draw_line((40, 130), (86, 130), color=BL, width=3)
    figtext(p, 63, 215, "sample + host", 13)
    arrow(p, 100, 110, 200, 105)
    figtext(p, 150, 80, "molten soft agar", 13)
    p.draw_oval(pymupdf.Rect(210, 68, 430, 212), fill=(0.06, 0.08, 0.11), color=BL, width=4)
    p.draw_oval(pymupdf.Rect(222, 80, 418, 200), fill=LAWN, color=None)
    for (x, y, r) in [(290, 120, 7), (330, 105, 6), (355, 140, 8), (300, 160, 6), (340, 168, 5), (270, 145, 5)]:
        p.draw_circle((x, y), r, fill=CLEAR, color=None)
    figtext(p, 320, 238, "even lawn + countable plaques", 13, (0.49, 0.55, 0.63))
    savefig(d, "overlay")
    # cscl
    d, p = figpage(200, 300)
    p.draw_rect(pymupdf.Rect(70, 20, 130, 260), color=WH, width=4, radius=0.4)
    p.draw_rect(pymupdf.Rect(74, 26, 126, 254), fill=(0.55, 0.6, 0.7), color=None, radius=0.4)
    p.draw_rect(pymupdf.Rect(74, 130, 126, 142), fill=(0.49, 0.83, 0.98), color=None)
    arrow(p, 150, 136, 132, 136)
    figtext(p, 166, 142, "?", 26, MG)
    figtext(p, 100, 285, "density gradient", 13, (0.49, 0.55, 0.63))
    savefig(d, "cscl")
    # morphotypes
    d, p = figpage(560, 240)
    p.draw_polyline([(90, 12), (130, 30), (130, 70), (90, 88), (50, 70), (50, 30)], color=WH, width=3.5, closePath=True)
    p.draw_rect(pymupdf.Rect(79, 88, 101, 148), color=WH, width=3.5)
    for a, b, c, e in [(74, 148, 56, 178), (106, 148, 124, 178), (82, 150, 78, 180), (98, 150, 102, 180)]:
        p.draw_line((a, b), (c, e), color=WH, width=3)
    figtext(p, 90, 200, "1", 20, (0.49, 0.83, 0.98))
    p.draw_polyline([(280, 12), (316, 28), (316, 64), (280, 80), (244, 64), (244, 28)], color=WH, width=3.5, closePath=True)
    p.draw_polyline([(280, 80), (266, 110), (294, 130), (280, 160), (270, 178), (280, 198)], color=WH, width=3.5)
    p.draw_line((280, 198), (266, 216), color=WH, width=3)
    p.draw_line((280, 198), (294, 216), color=WH, width=3)
    arrow(p, 320, 120, 294, 128)
    figtext(p, 338, 126, "?", 26, MG)
    figtext(p, 280, 232, "2", 20, (0.49, 0.83, 0.98))
    p.draw_polyline([(470, 12), (504, 28), (504, 64), (470, 80), (436, 64), (436, 28)], color=WH, width=3.5, closePath=True)
    p.draw_rect(pymupdf.Rect(461, 80, 479, 102), color=WH, width=3.5)
    for a, b, c, e in [(460, 102, 450, 120), (480, 102, 490, 120), (470, 102, 470, 122)]:
        p.draw_line((a, b), (c, e), color=WH, width=3)
    figtext(p, 470, 200, "3", 20, (0.49, 0.83, 0.98))
    savefig(d, "morphotypes")
    # biofilm
    d, p = figpage(460, 250)
    p.draw_line((30, 220), (430, 220), color=WH, width=5)
    for x0, h in [(40, 70), (120, 100), (220, 80), (320, 60)]:
        p.draw_bezier((x0, 220), (x0 + 20, 220 - h), (x0 + 70, 220 - h), (x0 + 90, 220), color=BL, width=3)
    for (x, y, r) in [(75, 196, 9), (92, 182, 8), (165, 170, 9), (185, 188, 8), (265, 180, 9), (283, 196, 8), (360, 192, 9)]:
        p.draw_circle((x, y), r, fill=(0.49, 0.83, 0.98), color=None)
    for y in (120, 90, 60):
        p.draw_line((40 + (120 - y), y), (420 - (120 - y), y), color=(0.23, 0.26, 0.34), width=2, dashes="4 6")
    arrow(p, 430, 70, 350, 96)
    figtext(p, 442, 64, "?", 26, MG)
    figtext(p, 230, 244, "cells embedded in a secreted matrix on a surface", 12, (0.49, 0.55, 0.63))
    savefig(d, "biofilm")
    # spot assay
    d, p = figpage(380, 280)
    p.draw_circle((190, 130), 120, fill=(0.06, 0.08, 0.11), color=BL, width=4)
    p.draw_circle((190, 130), 108, fill=LAWN, color=None)
    for (x, y, r) in [(140, 90, 16), (200, 80, 15), (250, 110, 14), (150, 160, 15), (215, 170, 13)]:
        p.draw_circle((x, y), r, fill=CLEAR, color=None)
    arrow(p, 300, 236, 262, 186)
    figtext(p, 316, 252, "?", 26, MG)
    figtext(p, 190, 272, "drops of phage stocks on a lawn", 13, (0.49, 0.55, 0.63))
    savefig(d, "spot_assay")
    # plaques turbid
    d, p = figpage(460, 240)
    p.draw_circle((120, 110), 92, fill=(0.06, 0.08, 0.11), color=BL, width=4)
    p.draw_circle((120, 110), 82, fill=LAWN, color=None)
    for (x, y, r) in [(95, 85, 10), (140, 75, 9), (155, 120, 10), (100, 135, 9), (130, 150, 8)]:
        p.draw_circle((x, y), r, fill=CLEAR, color=None)
    figtext(p, 120, 225, "Plate A", 17, (0.49, 0.83, 0.98))
    p.draw_circle((340, 110), 92, fill=(0.06, 0.08, 0.11), color=BL, width=4)
    p.draw_circle((340, 110), 82, fill=LAWN, color=None)
    for (x, y, r) in [(315, 85, 10), (360, 75, 9), (375, 120, 10), (320, 135, 9), (350, 150, 8)]:
        p.draw_circle((x, y), r, fill=(0.94, 0.90, 0.81), color=None)
        p.draw_circle((x, y), r * 0.4, fill=(0.80, 0.72, 0.56), color=None)
    figtext(p, 340, 225, "Plate B", 17, (0.49, 0.83, 0.98))
    savefig(d, "plaques_turbid")

make_figures()
FIGW = {n: pymupdf.Pixmap(str(FIGDIR / f"{n}.png")) for n in
        ["growth_curve", "plaque_halo", "t4_diagram", "lysogeny", "overlay", "cscl",
         "morphotypes", "biofilm", "spot_assay", "plaques_turbid"]}

OPT_RECTS = [(122, 291, 454, 351), (506, 291, 838, 351), (122, 384, 454, 444), (506, 384, 838, 444)]

def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt

def q_slides(r, q):
    for show in (False, True):
        s = slide()
        tb(s, 0.9, 0.5, 4, 0.4, [[(f"ROUND {r['num']}", 14, WHITE, True)]])
        if q.get("audience"):
            shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.6), 0.42, Inches(1.9), 0.4)
            shp.shadow.inherit = False
            shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(0xE6, 0x1A, 0x80)
            shp.line.fill.background()
            tf = shp.text_frame; p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
            rr = p0.add_run(); rr.text = "AUDIENCE"
            rr.font.size = Pt(13); rr.font.bold = True; rr.font.color.rgb = WHITE; rr.font.name = "Arial"
        if r.get("badge"):
            shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.3), 0.44, Inches(1.9), 0.36)
            shp.shadow.inherit = False
            shp.fill.solid(); shp.fill.fore_color.rgb = G2
            shp.line.fill.background()
            tf = shp.text_frame; p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
            rr = p0.add_run(); rr.text = r["badge"]
            rr.font.size = Pt(11); rr.font.bold = True; rr.font.color.rgb = WHITE; rr.font.name = "Arial"
        qy, qsize = 1.6, 24
        if q.get("figure"):
            pm = FIGW[q["figure"]]
            h_in = 2.35
            w_in = h_in * pm.width / pm.height
            s.shapes.add_picture(str(FIGDIR / f"{q['figure']}.png"),
                                 Inches((13.333 - w_in) / 2), Inches(0.85), Inches(w_in), Inches(h_in))
            qy, qsize = 3.3, (18 if len(q["q"]) <= 70 else 16)
        else:
            qsize = 24 if len(q["q"]) <= 70 else (20 if len(q["q"]) <= 110 else 18)
        tb(s, 1.2, qy, 10.93, 1.25, [[(q["q"], qsize, WHITE, True)]], align=PP_ALIGN.CENTER)
        for i, opt in enumerate(q["options"]):
            x0, y0, x1, y1 = OPT_RECTS[i]
            win = show and i == q["answer"]
            osize = 16 if len(opt) <= 40 else (14 if len(opt) <= 64 else 12)
            hexcard(s, x0 * PT, y0 * PT, (x1 - x0) * PT, (y1 - y0) * PT, "ABCD"[i], opt, win=win, size=osize)
        if not show:
            notes(s, f"ANSWER: {'ABCD'[q['answer']]} — {q['options'][q['answer']]}")
    return

def divider(r):
    s = slide()
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.15), 1.5, Inches(1.05), Inches(1.05))
    shp.shadow.inherit = False
    f = shp.fill; f.gradient()
    f.gradient_stops[0].color.rgb = G1; f.gradient_stops[1].color.rgb = G2
    shp.line.fill.background()
    tf = shp.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
    rr = p0.add_run(); rr.text = str(r["num"])
    rr.font.size = Pt(40); rr.font.bold = True; rr.font.color.rgb = WHITE; rr.font.name = "Arial"
    tb(s, 1, 3.0, 11.33, 0.8, [[(r["title"], 32, WHITE, True)]], align=PP_ALIGN.CENTER)
    if r.get("badge"):
        tb(s, 1, 3.75, 11.33, 0.4, [[("✦ " + r["badge"] + " ✦", 16, ACC, True)]], align=PP_ALIGN.CENTER)
    sc = r.get("scoring", {})
    chips = [f"Correct +{sc.get('correct', 0)}", f"Wrong {sc.get('wrong', 0)}", f"{sc.get('time', 30)} sec"]
    if sc.get("pass"):
        chips.append(f"passed +{sc['pass']}")
    widths = [0.16 * len(c) + 0.5 for c in chips]
    x = (13.333 - sum(widths) - 0.2 * (len(chips) - 1)) / 2
    for c, w in zip(chips, widths):
        chip(s, x, 4.5, c)
        x += w + 0.2

# ------------------------------- deck -------------------------------
s = slide()
tb(s, 1, 2.2, 11.33, 1.0, [[("START", 54, ACC, True)]], align=PP_ALIGN.CENTER)
tb(s, 1, 3.3, 11.33, 1.0, [[(data["meta"]["subtitle"], 28, WHITE, True)]], align=PP_ALIGN.CENTER)
tb(s, 1, 4.5, 11.33, 0.5, [[(data["meta"]["tagline"], 16, MUT, True)]], align=PP_ALIGN.CENTER)

s = slide()
tb(s, 1, 0.7, 11.33, 0.7, [[("Teams", 32, WHITE, True)]], align=PP_ALIGN.CENTER)
for i, t in enumerate(data["teams"]):
    x = 0.62 + i * 3.08
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.7), Inches(2.85), Inches(3.0))
    shp.shadow.inherit = False
    shp.fill.solid(); shp.fill.fore_color.rgb = CARD; shp.line.color.rgb = LINE
    tb(s, x + 0.15, 2.0, 2.55, 0.8, [[(t["name"], 17, SKY, True)]], align=PP_ALIGN.CENTER)
    tb(s, x + 0.15, 2.9, 2.55, 1.6, [[(mm, 13, MUT, True)] for mm in t["members"]], align=PP_ALIGN.CENTER)

s = slide()
tb(s, 1, 0.6, 11.33, 0.7, [[("Rounds", 32, WHITE, True)]], align=PP_ALIGN.CENTER)
paras = []
for r in data["rounds"]:
    label = f"BREAK ({r['minutes']} min)" if r["type"] == "break" else f"{r['num']}.  {r['title']}"
    line = [(label, 16, WHITE, True)]
    if r.get("badge"):
        line.append(("   " + r["badge"], 11, ACC, True))
    paras.append(line)
tb(s, 1.4, 1.5, 5.2, 5.4, paras[:7])
tb(s, 6.9, 1.5, 5.2, 5.4, paras[7:])

s = slide()
tb(s, 1, 0.7, 11.33, 0.7, [[("Markings", 32, WHITE, True)]], align=PP_ALIGN.CENTER)
rows = [("Rounds 1-5", "+10 / 0", "30 sec"), ("Rounds 6-10", "+10 / -5", "30 sec"),
        ("Rapid Fire", "+5 / 0, no passing", "60 sec / team"),
        ("Tie-Breaker", "+20 / -10, passed +5", "60 sec"),
        ("Emoji Rebus & Lytic Laughs", "+10 / 0", "30 sec")]
paras = [[(a, 18, SKY, True), ("      " + b, 18, WHITE, True), ("      " + c, 15, MUT, True)] for a, b, c in rows]
paras.append([("Each MCQ round: 8 team questions + 2 audience questions", 14, MUT, True)])
tb(s, 2.2, 1.8, 9, 4.5, paras)

for r in data["rounds"]:
    if r["type"] == "break":
        s = slide()
        tb(s, 1, 2.6, 11.33, 1.0, [[("Break Time!", 44, WHITE, True)]], align=PP_ALIGN.CENTER)
        tb(s, 1, 3.8, 11.33, 0.7, [[(f"{r['minutes']} MIN", 28, ACC, True)]], align=PP_ALIGN.CENTER)
        continue
    divider(r)
    if r["type"] == "rapid":
        for tset in r["teams"]:
            tname = data["teams"][tset["team"]]["name"]
            for qi, qq in enumerate(tset["questions"]):
                for show in (False, True):
                    s = slide()
                    tb(s, 0.9, 0.5, 6, 0.4, [[(f"ROUND {r['num']} — RAPID FIRE", 14, WHITE, True)]])
                    tb(s, 1, 1.3, 11.33, 0.4, [[(f"{tname}  ·  question {qi + 1} of {len(tset['questions'])}", 15, MUT, True)]], align=PP_ALIGN.CENTER)
                    tb(s, 1.2, 2.4, 10.93, 1.2, [[(qq["q"], 26 if len(qq["q"]) <= 70 else 22, WHITE, True)]], align=PP_ALIGN.CENTER)
                    if show:
                        hexcard(s, 3.0, 4.4, 7.33, 0.95, "", qq["a"], win=True, size=20 if len(qq["a"]) <= 40 else 16)
                    else:
                        hexcard(s, 3.0, 4.4, 7.33, 0.95, "", "?", win=False, size=20)
                        notes(s, f"ANSWER: {qq['a']}")
        continue
    for q in r["questions"]:
        q_slides(r, q)

s = slide()
tb(s, 1, 3.0, 11.33, 1.0, [[("Thank you for your answers!", 38, WHITE, True)]], align=PP_ALIGN.CENTER)
tb(s, 1, 4.1, 11.33, 0.5, [[("See the scoreboard for final standings", 16, MUT, True)]], align=PP_ALIGN.CENTER)

out = HERE / "phage-quiz.pptx"
prs.save(str(out))
print("wrote", out, len(prs.slides._sldIdLst), "slides")
